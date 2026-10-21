from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .settings import settings


def _sheet(wb: Workbook, title: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title=title)
    if not rows:
        ws.append(["No data"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9EAF7")
        c.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    for i, header in enumerate(headers, 1):
        max_len = max(len(str(header)), *(len(str(ws.cell(r, i).value or "")) for r in range(2, ws.max_row + 1)))
        ws.column_dimensions[get_column_letter(i)].width = min(max(max_len + 2, 12), 44)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def export_result(job_id: str, result: dict[str, Any]) -> Path:
    wb = Workbook()
    wb.remove(wb.active)

    summary = result.get("summary", {})
    _sheet(wb, "Run Summary", [{"metric": k, "value": v} for k, v in summary.items()])
    _sheet(wb, "Device Summary", list(result.get("device_summary") or []))
    _sheet(wb, "Channel Analysis", list(result.get("channel_analysis") or []))
    _sheet(wb, "Video Gaps", list(result.get("video_gaps") or []))
    _sheet(wb, "Engine Events", list(result.get("engine_events") or []))
    _sheet(wb, "Engine Intervals", list(result.get("engine_intervals") or []))
    _sheet(wb, "Engine Unverified Windows", list(result.get("engine_unverified_intervals") or []))
    _sheet(wb, "Video File List", list(result.get("video_files") or []))

    collection = result.get("collection") or {}
    collection_rows = []
    for key, value in collection.items():
        if isinstance(value, (list, dict)):
            continue
        collection_rows.append({"item": key, "value": value})
    engine_diag = result.get("engine_diagnostics") or {}
    vehicle_mapping = engine_diag.get("vehicle_mapping") or {}
    mapping_dict = vehicle_mapping.get("mapping") if isinstance(vehicle_mapping, dict) else {}
    if isinstance(mapping_dict, dict):
        _sheet(wb, "BSJ Vehicle Map", [
            {"imei": imei, "bsj_vehicle_id": vehicle_id}
            for imei, vehicle_id in sorted(mapping_dict.items())
        ])
    for key, value in engine_diag.items():
        if isinstance(value, (list, dict)):
            continue
        collection_rows.append({"item": f"engine_{key}", "value": value})
    _sheet(wb, "Collection Info", collection_rows)
    _sheet(wb, "Wake API Failures", list(collection.get("wake_failures") or []))
    _sheet(wb, "File List API Failures", list(collection.get("file_list_failures") or []))
    _sheet(wb, "Vehicle Mapping Failures", [{"imei": imei} for imei in (collection.get("vehicle_mapping_failed") or [])])
    _sheet(wb, "Engine Status API Failures", list(collection.get("engine_status_failures") or []))

    path = settings.export_dir / f"bsj_recording_completeness_full_api_{job_id[:8]}.xlsx"
    wb.save(path)
    return path
